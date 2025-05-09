import streamlit as st
from docx import Document
import os
import tempfile

MODULE_FORMATS = {
    "AUDIT_TOOLS": ["docx"],
    "NETWORK_DIAGRAM": ["jpg", "png"],
    "DECISION_MATRIX_TABLE": ["jpg", "xls", "xlsx", "docx"],
    "APP_MIG_PLAN": ["docx"],
    "DR_METHOD": ["docx"],
    "SLA_APPS": ["docx", "xls", "xlsx", "jpg", "png"],
    "DC_NB": ["docx"]
}

st.set_page_config(page_title="RFP Modular Builder", layout="wide")
st.title("📄 RFP Modular Builder + Uploads + Placeholder Input")

TEMPLATE_DIR = "templates"
os.makedirs(TEMPLATE_DIR, exist_ok=True)

uploaded_template = st.sidebar.file_uploader("📄 Upload Word Template", type=["docx"], key="main_docx_template")
if uploaded_template:
    with open(os.path.join(TEMPLATE_DIR, uploaded_template.name), "wb") as f:
        f.write(uploaded_template.read())
    st.sidebar.success(f"✅ Template {uploaded_template.name} uploaded.")

template_files = [f for f in os.listdir(TEMPLATE_DIR) if f.endswith(".docx")]
selected_template = st.sidebar.selectbox("📌 Choose a Template", template_files)

customer_input_file = st.sidebar.file_uploader("📥 Upload customer_inputs.txt", type=["txt"])
st.sidebar.header("📌 Manual Inputs")
manual_placeholders = {
    "CUSTOMER_NAME": st.sidebar.text_input("Customer Name", "Telco ABC"),
    "COMPANY_NAME": st.sidebar.text_input("Company Name", "Rackspace"),
    "CITY_NAME": st.sidebar.text_input("City", "Hanoi"),
    "PARTNER_NAME": st.sidebar.text_input("Partner Name", "CMC")
}

placeholder_data = manual_placeholders.copy()
if customer_input_file:
    content = customer_input_file.read().decode("utf-8")
    for line in content.splitlines():
        if "=" in line:
            key, val = line.split("=", 1)
            placeholder_data[key.strip()] = val.strip().strip('"')

st.sidebar.header("📂 Choose Modules")
module_dirs = sorted([d for d in os.listdir("modules") if os.path.isdir(f"modules/{d}")])
selected_modules = st.sidebar.multiselect("Modules to include", module_dirs, default=module_dirs)

st.sidebar.header("📁 Upload or Edit Modules")
uploaded_assets = {}
edited_modules = {}

with st.form("module_editor"):
    for module in selected_modules:
        module_path = f"modules/{module}/main.txt"
        os.makedirs(os.path.dirname(module_path), exist_ok=True)
        accepted_types = MODULE_FORMATS.get(module.upper(), ["txt"])
        uploaded = st.file_uploader(f"Upload for {module}", type=accepted_types, key=f"{module}_upload")
        if uploaded:
            uploaded_assets[module] = uploaded
            content = f"[See uploaded file: {uploaded.name}]"
        else:
            if os.path.exists(module_path):
                with open(module_path) as f:
                    content = f.read()
            else:
                content = ""
            for k, v in placeholder_data.items():
                content = content.replace(f"{{{k}}}", v)
        updated_text = st.text_area(f"✏️ Edit {module}", value=content, height=250, key=module)
        edited_modules[module] = updated_text
    save_btn = st.form_submit_button("💾 Save Modules")

if save_btn:
    for mod, text in edited_modules.items():
        with open(f"modules/{mod}/main.txt", "w") as f:
            f.write(text)
    st.success("✅ Modules saved.")

if selected_template and st.sidebar.button("📄 Generate RFP"):
    doc = Document(os.path.join(TEMPLATE_DIR, selected_template))
    for para in doc.paragraphs:
        for mod, txt in edited_modules.items():
            if f"{{{mod}}}" in para.text:
                para.text = para.text.replace(f"{{{mod}}}", txt)
        for k, v in placeholder_data.items():
            para.text = para.text.replace(f"{{{k}}}", v)

    tmp_path = tempfile.NamedTemporaryFile(delete=False, suffix=".docx").name
    doc.save(tmp_path)
    with open(tmp_path, "rb") as f:
        st.download_button("📥 Download Final RFP", f, "Generated_RFP.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")